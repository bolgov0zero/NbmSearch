import gc
import os
import time
import heapq
import logging
import threading
from concurrent.futures import ThreadPoolExecutor

from watchdog.observers import Observer
from watchdog.events import FileSystemEventHandler

from app.settings import MAX_WORKERS, HEAVY_WORKERS, MAX_FILE_BYTES, MAX_TEXT_CHARS
from app import database as db

logger = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS = {".docx", ".doc", ".xlsx", ".xls", ".pdf", ".rtf", ".txt", ".csv"}

# Formats that load entire file into memory — limit their concurrency separately
_HEAVY_EXTENSIONS = {".pdf", ".xlsx", ".xls"}
_heavy_sem = threading.Semaphore(HEAVY_WORKERS)

# ── Watchdog event log ────────────────────────────────────────────────────────
# folder_id → list of {"ts": float, "type": str, "name": str, "path": str}
_LOG_TTL = 7 * 24 * 3600   # 7 days in seconds
_LOG_MAX = 1000             # max entries per folder (oldest pruned first)

_watchdog_log: dict[int, list] = {}
_watchdog_log_lock = threading.Lock()


def _log_event(folder_id: int, event_type: str, path: str):
    name = os.path.basename(path)
    entry = {"ts": time.time(), "type": event_type, "name": name, "path": path}
    with _watchdog_log_lock:
        log = _watchdog_log.setdefault(folder_id, [])
        log.insert(0, entry)
        # Prune entries older than TTL and keep max size
        cutoff = time.time() - _LOG_TTL
        _watchdog_log[folder_id] = [e for e in log if e["ts"] >= cutoff][:_LOG_MAX]


def get_watchdog_log(folder_id: int) -> list:
    cutoff = time.time() - _LOG_TTL
    with _watchdog_log_lock:
        return [e for e in _watchdog_log.get(folder_id, []) if e["ts"] >= cutoff]


def clear_watchdog_log(folder_id: int):
    with _watchdog_log_lock:
        _watchdog_log[folder_id] = []


def get_created_since(since: float) -> dict[int, int]:
    """folder_id → number of 'created' (new file) events logged after `since`."""
    out: dict[int, int] = {}
    with _watchdog_log_lock:
        for fid, log in _watchdog_log.items():
            n = sum(1 for e in log if e["ts"] > since and e["type"] == "created")
            if n:
                out[fid] = n
    return out


# ── Progress tracking ─────────────────────────────────────────────────────────
# folder_id → {"total": int, "done": int, "status": "idle"|"indexing"|"done"}
_progress: dict[int, dict] = {}
_progress_lock = threading.Lock()


def get_progress(folder_id: int) -> dict:
    with _progress_lock:
        return dict(_progress.get(folder_id, {"total": 0, "done": 0, "status": "idle"}))


def _set_progress(folder_id: int, total: int, done: int, status: str):
    with _progress_lock:
        _progress[folder_id] = {"total": total, "done": done, "status": status}


# ── Reindex scheduler ─────────────────────────────────────────────────────────

def reindex_scheduler():
    while True:
        time.sleep(30)
        try:
            now = time.time()
            schedules = db.get_schedules()
            folders = {f["id"]: f for f in db.get_folders()}
            for s in schedules:
                nxt = s.get("next_run_at") or 0
                if now < nxt:
                    continue
                folder = folders.get(s["folder_id"])
                if not folder:
                    continue
                db.update_schedule_run(s["id"], s["reindex_minutes"])
                if is_indexing(folder["id"]):
                    logger.info("Scheduled reindex of '%s' skipped: previous run is still in progress",
                                folder["name"])
                    continue
                threading.Thread(target=index_folder, args=(folder,), daemon=True).start()
        except Exception as e:  # поток планировщика не должен умирать из-за ошибки БД
            logger.error("Reindex scheduler error: %s", e)


# ── Text extraction ───────────────────────────────────────────────────────────

def _read_text_file(path: str) -> str:
    # Читаем не больше MAX_TEXT_CHARS: гигабайтный csv/txt не должен целиком попадать в память
    with open(path, "r", encoding="utf-8", errors="ignore") as f:
        return f.read(MAX_TEXT_CHARS)


def extract_text(path: str) -> str:
    ext = os.path.splitext(path)[1].lower()
    try:
        # txt/csv читаются потоково с обрезкой; остальные форматы грузятся целиком —
        # слишком большие файлы пропускаем (в индекс попадёт только имя файла)
        if MAX_FILE_BYTES and ext not in (".txt", ".csv"):
            size = os.path.getsize(path)
            if size > MAX_FILE_BYTES:
                logger.warning("File too large (%.1f MB > %.1f MB), content skipped: %s",
                               size / (1024 * 1024), MAX_FILE_BYTES / (1024 * 1024), path)
                return ""

        if ext in (".txt", ".csv"):
            return _read_text_file(path)

        elif ext == ".docx":
            import docx
            doc = docx.Document(path)
            return "\n".join(p.text for p in doc.paragraphs)[:MAX_TEXT_CHARS]

        elif ext == ".doc":
            return _extract_doc(path)[:MAX_TEXT_CHARS]

        elif ext == ".xlsx":
            import openpyxl
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            try:
                parts = []
                total = 0
                for ws in wb.worksheets:
                    if total >= MAX_TEXT_CHARS:
                        break
                    try:
                        for row in ws.iter_rows(values_only=True):
                            try:
                                line = " ".join(str(c) for c in row if c is not None)
                            except Exception:
                                continue
                            parts.append(line)
                            total += len(line) + 1
                            if total >= MAX_TEXT_CHARS:
                                break
                    except Exception:
                        pass
                return "\n".join(parts)
            finally:
                wb.close()

        elif ext == ".xls":
            import xlrd
            wb = xlrd.open_workbook(path)
            try:
                parts = []
                total = 0
                for sheet in wb.sheets():
                    if total >= MAX_TEXT_CHARS:
                        break
                    for row_idx in range(sheet.nrows):
                        line = " ".join(str(sheet.cell_value(row_idx, c))
                                        for c in range(sheet.ncols))
                        parts.append(line)
                        total += len(line) + 1
                        if total >= MAX_TEXT_CHARS:
                            break
                return "\n".join(parts)
            finally:
                wb.release_resources()

        elif ext == ".pdf":
            import pdfplumber
            with pdfplumber.open(path) as pdf:
                parts = []
                total = 0
                for page in pdf.pages:
                    try:
                        text = page.extract_text() or ""
                    finally:
                        # pdfplumber кэширует разобранные объекты страницы; без сброса
                        # память растёт пропорционально размеру документа
                        flush = getattr(page, "flush_cache", None)
                        if flush:
                            flush()
                    parts.append(text)
                    total += len(text) + 1
                    if total >= MAX_TEXT_CHARS:
                        break
                return "\n".join(parts)

        elif ext == ".rtf":
            from striprtf.striprtf import rtf_to_text
            with open(path, "r", encoding="utf-8", errors="ignore") as f:
                return rtf_to_text(f.read())[:MAX_TEXT_CHARS]

    except Exception as e:
        logger.warning("Cannot extract text from %s: %s", path, e)
    return ""


def _extract_doc(path: str) -> str:
    import re
    import olefile
    try:
        with olefile.OleFileIO(path) as ole:
            if not ole.exists("WordDocument"):
                return ""
            stream = ole.openstream("WordDocument").read()
            text = stream.decode("utf-16-le", errors="ignore")
            text = re.sub(r"[^\x20-\x7eЀ-ӿ\n\r\t]+", " ", text)
            return " ".join(text.split())
    except Exception:
        with open(path, "rb") as f:
            raw = f.read()
        import re
        text = raw.decode("utf-16-le", errors="ignore")
        text = re.sub(r"[^\x20-\x7eЀ-ӿ\n\r\t]+", " ", text)
        return " ".join(text.split())


# ── Indexing ──────────────────────────────────────────────────────────────────

def index_file(folder_id: int, path: str) -> bool:
    """Returns True if the file was newly added to the index, False otherwise."""
    ext = os.path.splitext(path)[1].lower()
    if ext not in SUPPORTED_EXTENSIONS:
        return False
    if os.path.basename(path).startswith("~$"):
        return False
    try:
        stat = os.stat(path)
        if stat.st_size == 0:
            return False  # skip empty files (being written / temp placeholders)
        if ext in _HEAVY_EXTENSIONS:
            with _heavy_sem:
                content = extract_text(path)
        else:
            content = extract_text(path)
        name = os.path.basename(path)
        is_new = db.upsert_file(
            folder_id=folder_id,
            path=path,
            name=name,
            size=stat.st_size,
            modified_at=stat.st_mtime,
            content=f"{name}\n{content}",
            indexed_at=time.time(),
        )
        logger.debug("Indexed: %s", path)
        return is_new
    except FileNotFoundError:
        return False
    except Exception as e:
        logger.error("Error indexing %s: %s", path, e)
        return False


_running_folders: set[int] = set()
_running_lock = threading.Lock()


def is_indexing(folder_id: int | None = None) -> bool:
    """Идёт ли индексация указанной папки (или любой, если folder_id не задан)."""
    with _running_lock:
        return (folder_id in _running_folders) if folder_id is not None else bool(_running_folders)


def index_folder(folder: dict, full: bool = False, update_last_reindex: bool = True) -> bool:
    """Индексирует папку. Две индексации одной папки одновременно не запускаются
    (иначе они дублируют работу и кратно увеличивают потребление памяти).
    Возвращает False, если запуск пропущен, потому что папка уже индексируется."""
    folder_id = folder["id"]
    with _running_lock:
        if folder_id in _running_folders:
            logger.warning("Folder '%s' is already being indexed, duplicate run skipped", folder["name"])
            return False
        _running_folders.add(folder_id)
    try:
        _index_folder(folder, full, update_last_reindex)
        return True
    finally:
        with _running_lock:
            _running_folders.discard(folder_id)
        p = get_progress(folder_id)
        if p["status"] == "indexing":  # прервано ошибкой — не оставляем прогресс-бар «висеть»
            _set_progress(folder_id, p["total"], p["done"], "done")
        gc.collect()


def _index_folder(folder: dict, full: bool, update_last_reindex: bool):
    folder_id = folder["id"]
    folder_path = folder["path"]
    folder_name = folder["name"]
    last_reindex_at = None if full else folder.get("last_reindex_at")  # None → full index

    if not os.path.isdir(folder_path):
        logger.warning("Folder does not exist: %s", folder_path)
        return

    all_files = []
    for root, _, filenames in os.walk(folder_path):
        for fname in filenames:
            if os.path.splitext(fname)[1].lower() in SUPPORTED_EXTENSIONS and not fname.startswith("~$"):
                all_files.append(os.path.join(root, fname))

    # Get already-indexed paths: needed for stale detection and new-file detection
    indexed_paths = db.get_all_paths_in_folder(folder_id)

    if last_reindex_at:
        # Incremental: only new files or files modified after last reindex
        to_index = []
        for f in all_files:
            try:
                if os.stat(f).st_mtime > last_reindex_at or f not in indexed_paths:
                    to_index.append(f)
            except OSError:
                pass
        logger.info("Indexing folder '%s': %d changed/new of %d total (since %s)",
                    folder_name, len(to_index), len(all_files),
                    time.strftime("%d.%m.%Y %H:%M", time.localtime(last_reindex_at)))
    else:
        to_index = all_files
        logger.info("Indexing folder '%s': %d files (full)", folder_name, len(to_index))

    total = len(to_index)
    _set_progress(folder_id, total, 0, "indexing")

    done = 0
    done_lock = threading.Lock()

    def _index_and_count(path):
        nonlocal done
        index_file(folder_id, path)
        with done_lock:
            done += 1
            current = done
        _set_progress(folder_id, total, current, "indexing")
        if current % 500 == 0:
            gc.collect()  # парсеры xls/xlsx/pdf оставляют циклический мусор

    with ThreadPoolExecutor(max_workers=MAX_WORKERS) as executor:
        executor.map(_index_and_count, to_index)

    # Remove deleted files
    current_paths = set(all_files)
    stale_paths = indexed_paths - current_paths
    logger.info("Folder '%s' stats: on_disk=%d in_db=%d stale=%d",
                folder_name, len(all_files), len(indexed_paths), len(stale_paths))
    for stale in stale_paths:
        logger.info("Removing stale: %s", stale)
        db.delete_file_from_folder(folder_id, stale)

    _set_progress(folder_id, total, total, "done")
    if update_last_reindex:
        db.set_folder_last_reindex(folder_id, time.time())
    db.update_folder_file_count(folder_id)
    logger.info("Done indexing folder '%s'", folder_name)


# ── Watchdog ──────────────────────────────────────────────────────────────────

_observer: Observer | None = None
_observer_lock = threading.Lock()


_WATCHDOG_DEBOUNCE = 5.0  # seconds to wait after last event before indexing


_DELETE_DEBOUNCE = 5.0  # seconds to wait before confirming a deletion

_COUNT_DEBOUNCE = 3.0  # file_count в основной БД обновляется не чаще, чем раз в N секунд

# Разбор файлов по событиям — в небольшом пуле: при массовом копировании/сканировании
# шары событий могут быть сотни тысяч, потоков при этом должно быть фиксированное число.
WATCHDOG_WORKERS = max(1, MAX_WORKERS // 2)


class _EventDispatcher:
    """Debounce событий файловой системы + ограниченный пул воркеров.

    Заменяет схему «threading.Timer на каждое событие»: при шторме событий (бэкап,
    антивирус, массовое копирование) она создавала сотни тысяч потоков и разбирала
    файлы без ограничения параллелизма. Здесь отложенные задачи — это записи в словаре
    и куче, а выполняют их WATCHDOG_WORKERS потоков.
    """

    def __init__(self, workers: int):
        self._workers = workers
        self._cv = threading.Condition()
        self._heap: list[tuple[float, int, tuple]] = []   # (due, seq, key)
        self._items: dict[tuple, tuple] = {}              # key -> (seq, fn, args)
        self._seq = 0
        self._threads: list[threading.Thread] = []

    def _ensure_started(self):
        if self._threads:
            return
        for i in range(self._workers):
            t = threading.Thread(target=self._loop, name=f"nbm-watchdog-{i}", daemon=True)
            t.start()
            self._threads.append(t)

    def schedule(self, key: tuple, delay: float, fn, *args, replace: bool = True):
        """Выполнить fn(*args) через delay секунд. Повторный вызов с тем же key
        переносит срок (replace=True) либо игнорируется (replace=False)."""
        with self._cv:
            self._ensure_started()
            if not replace and key in self._items:
                return
            self._seq += 1
            self._items[key] = (self._seq, fn, args)
            heapq.heappush(self._heap, (time.monotonic() + delay, self._seq, key))
            self._cv.notify()

    def cancel(self, key: tuple) -> bool:
        with self._cv:
            return self._items.pop(key, None) is not None

    def clear(self):
        with self._cv:
            self._items.clear()
            self._heap.clear()

    def pending(self) -> int:
        with self._cv:
            return len(self._items)

    def _next_task(self):
        with self._cv:
            while True:
                if not self._heap:
                    self._cv.wait()
                    continue
                due, seq, key = self._heap[0]
                cur = self._items.get(key)
                if cur is None or cur[0] != seq:      # отменена или перенесена
                    heapq.heappop(self._heap)
                    continue
                wait = due - time.monotonic()
                if wait > 0:
                    self._cv.wait(wait)
                    continue
                heapq.heappop(self._heap)
                del self._items[key]
                return cur[1], cur[2]

    def _loop(self):
        while True:
            fn, args = self._next_task()
            try:
                fn(*args)
            except Exception as e:
                logger.error("Watchdog task failed: %s", e)


_dispatcher = _EventDispatcher(WATCHDOG_WORKERS)


def _schedule_count_update(folder_id: int):
    _dispatcher.schedule(("c", folder_id), _COUNT_DEBOUNCE,
                         db.update_folder_file_count, folder_id, replace=False)


def _is_temp(path: str) -> bool:
    name = os.path.basename(path)
    return (
        name.startswith("~$") or          # Office lock files
        name.startswith("._") or           # macOS resource forks
        name.startswith(".~") or           # LibreOffice temp
        name.endswith(".tmp") or           # Windows atomic save temps
        name.endswith(".TMP") or
        name.startswith("~") and name.endswith(".tmp")
    )


def _is_relevant(path: str) -> bool:
    """Событие интересно, только если это поддерживаемый формат и не временный файл."""
    return os.path.splitext(path)[1].lower() in SUPPORTED_EXTENSIONS and not _is_temp(path)


class FileEventHandler(FileSystemEventHandler):
    def __init__(self, folder_id: int):
        self.folder_id = folder_id
        self._event_types: dict[str, str] = {}   # path → first event type
        self._types_lock = threading.Lock()

    def _key(self, kind: str, path: str) -> tuple:
        return (kind, self.folder_id, path)

    def _schedule(self, path: str, event_type: str = "modified"):
        """Debounce: index the file only after DEBOUNCE seconds of silence.
        The first event type wins — created stays created even if modified follows."""
        with self._types_lock:
            self._event_types.setdefault(path, event_type)
        _dispatcher.schedule(self._key("i", path), _WATCHDOG_DEBOUNCE, self._do_index, path)

    def _do_index(self, path: str):
        with self._types_lock:
            event_type = self._event_types.pop(path, "modified")
        # Skip if file disappeared or is still empty (being written)
        try:
            stat = os.stat(path)
            if stat.st_size == 0:
                logger.debug("Watchdog: skipping empty file %s", path)
                return
        except OSError:
            return
        # Skip if mtime and size match what's already indexed (false positive event)
        meta = db.get_file_meta(self.folder_id, path)
        if meta is not None:
            db_mtime, db_size = meta
            if abs(stat.st_mtime - db_mtime) < 1.0 and stat.st_size == db_size:
                logger.debug("Watchdog: skipping unchanged file %s", path)
                return
        is_new = index_file(self.folder_id, path)
        _schedule_count_update(self.folder_id)
        if event_type == "created" and not is_new:
            event_type = "modified"
        _log_event(self.folder_id, event_type, path)

    @staticmethod
    def _is_temp(path: str) -> bool:
        return _is_temp(path)

    def _cancel_delete(self, path: str) -> bool:
        """Cancel pending delete. Returns True if one was cancelled."""
        return _dispatcher.cancel(self._key("d", path))

    def _do_delete(self, path: str):
        db.delete_file_from_folder(self.folder_id, path)
        _schedule_count_update(self.folder_id)
        _log_event(self.folder_id, "deleted", path)

    def _do_move_source(self, path: str):
        db.delete_file_from_folder(self.folder_id, path)
        _schedule_count_update(self.folder_id)

    def on_created(self, event):
        if event.is_directory or not _is_relevant(event.src_path):
            return
        # If file reappeared after a pending delete — it was an atomic save, log as created only
        self._cancel_delete(event.src_path)
        self._schedule(event.src_path, "created")

    def on_modified(self, event):
        if not event.is_directory and _is_relevant(event.src_path):
            self._cancel_delete(event.src_path)
            self._schedule(event.src_path, "modified")

    def on_deleted(self, event):
        if event.is_directory or not _is_relevant(event.src_path):
            return
        # Defer deletion — file may reappear immediately (atomic save pattern)
        _dispatcher.schedule(self._key("d", event.src_path), _DELETE_DEBOUNCE,
                             self._do_delete, event.src_path)

    def on_moved(self, event):
        if event.is_directory:
            return
        # Не блокируем поток наблюдателя обращением к БД — иначе при шторме событий
        # буфер ReadDirectoryChangesW переполняется и события теряются
        if _is_relevant(event.src_path):
            _dispatcher.schedule(self._key("m", event.src_path), 0.0,
                                 self._do_move_source, event.src_path)
        if _is_relevant(event.dest_path):
            self._schedule(event.dest_path, "moved")


def restart_watchdog():
    global _observer
    with _observer_lock:
        if _observer is not None:
            try:
                _observer.stop()
                _observer.join(timeout=5)
            except Exception:
                pass
            _observer = None
        _dispatcher.clear()   # события снятых с наблюдения папок больше не нужны

        folders = [f for f in db.get_folders() if f.get("watchdog_enabled")]
        if not folders:
            return

        observer = Observer()
        for folder in folders:
            if os.path.isdir(folder["path"]):
                handler = FileEventHandler(folder["id"])
                observer.schedule(handler, folder["path"], recursive=True)
                logger.info("Watchdog watching: %s", folder["path"])
        observer.start()
        _observer = observer
